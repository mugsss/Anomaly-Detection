"""Step 1 of the onboarding state machine: validate an uploaded loan tape.

Triggered by the Step Functions state machine after an S3 ObjectCreated event
under raw-uploads/{lender_id}/. Checks that the file is a readable Excel
workbook and that the columns the Part 1 pipeline depends on are present.
Does not run detection; that happens in the Fargate step. This is a cheap,
fast gate so a malformed export fails in seconds, not after a multi-minute
Fargate task spins up.

Raises on any problem. Step Functions catches the exception, retries with
backoff for transient S3 issues, and routes to the DLQ after retries are
exhausted (see loan_pipeline_stack.py).
"""

from __future__ import annotations

import io
import json
import logging
import os

import boto3
import openpyxl

logger = logging.getLogger()
logger.setLevel(logging.INFO)

s3 = boto3.client("s3")

# Columns the Part 1 loader treats as required. Kept in sync with
# src/loader.py's COLUMN_MAP; a mismatch here should be caught by CI, not
# discovered in production.
REQUIRED_COLUMNS = {
    "loan id",
    "borrower id",
    "loan amount",
    "interest rate",
    "loan term",
    "disbursal date",
    "payments",
}


def _normalise(name: object) -> str:
    return str(name).strip().lower()


def handler(event: dict, context) -> dict:
    """Validate the uploaded object referenced in the Step Functions input.

    Expects `event` to carry the S3 bucket and key, either directly (when
    invoked from a test) or nested under `detail` (the shape of an S3
    ObjectCreated event delivered via EventBridge).
    """
    detail = event.get("detail", event)
    bucket = detail["bucket"]["name"]
    key = detail["object"]["key"]

    logger.info("Validating s3://%s/%s", bucket, key)

    lender_id = _lender_id_from_key(key)
    response = s3.get_object(Bucket=bucket, Key=key)
    body = response["Body"].read()

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"file is not a readable Excel workbook: {exc}") from exc

    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise ValueError("workbook has no header row")

    present = {_normalise(c) for c in header_row if c is not None}
    missing = REQUIRED_COLUMNS - present
    if missing:
        raise ValueError(f"missing required columns: {sorted(missing)}")

    row_count = sum(1 for _ in sheet.iter_rows(min_row=2)) if sheet.max_row else 0
    logger.info("Validation passed: %d data rows, lender=%s", row_count, lender_id)

    return {
        "bucket": bucket,
        "key": key,
        "lender_id": lender_id,
        "row_count": row_count,
    }


def _lender_id_from_key(key: str) -> str:
    """Extract the lender id from a raw-uploads/{lender_id}/{file} key."""
    parts = key.split("/")
    if len(parts) < 3 or parts[0] != "raw-uploads":
        raise ValueError(f"key does not follow raw-uploads/<lender_id>/<file> layout: {key}")
    return parts[1]
